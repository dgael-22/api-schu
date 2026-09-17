# -*- coding: utf-8 -*-
"""
processors/exporter.py
======================
Escribe los dos archivos de salida:

  output/products_export_limpio.xlsx   formato Shopify, misma estructura y
                                       mismo orden de columnas que la plantilla
  output/reporte_proceso.xlsx          resumen y detalle de todo lo detectado

Reglas de formato de Shopify que se respetan aquí:
  * La PRIMERA fila de cada Handle lleva los datos del producto, la primera
    variante y la primera imagen.
  * Las filas siguientes llevan sólo Handle + campos de variante (+ imagen
    si quedan imágenes por colocar).
  * Si hay más imágenes que variantes, se añaden filas con sólo Handle,
    Image Src e Image Position.
  * SKU, código de barras, tallas y handle se escriben como TEXTO, para que
    Excel no se coma los ceros a la izquierda ni use notación científica.
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import rules
from processors.products import Producto
from processors.validator import Hallazgo

log = logging.getLogger("procesador.exportador")

_AZUL = PatternFill("solid", fgColor="1F3864")
_BLANCO_NEGRITA = Font(color="FFFFFF", bold=True)


# ---------------------------------------------------------------------------
# Excel de productos
# ---------------------------------------------------------------------------

def _celda_texto(valor: str) -> str:
    """Prefija con apóstrofo sólo si así se pidió en la configuración."""
    if valor and rules.SKU_PREFIJO_APOSTROFO:
        return f"'{valor}"
    return valor


def columnas_filtros() -> list[str]:
    """Encabezados de los metafields de filtro, en el formato que Shopify importa."""
    if not getattr(rules, "METAFIELDS_EN_CSV", False):
        return []
    ns = rules.METAFIELDS_NAMESPACE
    return [f"{nombre} (product.metafields.{ns}.{clave})"
            for clave, (nombre, _clases) in rules.METAFIELDS_FILTRO.items()]


def columnas_con_filtros(columnas: list[str]) -> list[str]:
    """Las de la plantilla, en su orden, y al final las de filtros que falten."""
    return list(columnas) + [c for c in columnas_filtros() if c not in columnas]


def valores_filtros(producto: Producto) -> dict[str, str]:
    """
    Valores de cada columna de filtro para un producto. Salen de sus tags, que
    ya pasaron por el vocabulario: clasificarlos otra vez da las mismas clases.
    """
    if not getattr(rules, "METAFIELDS_EN_CSV", False):
        return {}
    from processors import vocabulario

    tags = [t.strip() for t in producto.tags.split(",") if t.strip()]
    por_clave = vocabulario.metafields(vocabulario.clasificar(tags, producto.titulo))
    ns = rules.METAFIELDS_NAMESPACE
    return {f"{nombre} (product.metafields.{ns}.{clave})":
            rules.METAFIELDS_SEPARADOR_CSV.join(por_clave.get(clave, []))
            for clave, (nombre, _clases) in rules.METAFIELDS_FILTRO.items()}


def construir_filas(productos: list[Producto], columnas: list[str],
                    inventario: str | None = None,
                    inventario_rastreado: bool | None = None) -> pd.DataFrame:
    """
    Convierte los productos al layout de filas de Shopify.

    inventario            cantidad que va en "Variant Inventory Qty".
                          Por defecto, rules.INVENTARIO_POR_DEFECTO.
    inventario_rastreado  si Shopify debe llevar la cuenta. Cuando es False, la
                          columna "Variant Inventory Tracker" se deja vacía y
                          Shopify no descuenta al vender: la cantidad no cambia.
                          Por defecto, rules.INVENTARIO_RASTREADO.

    Los dos se reciben como parámetro (en vez de leerse de la configuración
    global) para que la API pueda procesar dos catálogos a la vez con ajustes
    distintos sin que uno pise al otro.
    """
    filas: list[dict] = []
    defaults = rules.DEFAULTS_SHOPIFY

    cantidad = str(rules.INVENTARIO_POR_DEFECTO if inventario is None else inventario)
    rastreado = (rules.INVENTARIO_RASTREADO if inventario_rastreado is None
                 else bool(inventario_rastreado))
    rastreador = defaults["Variant Inventory Tracker"] if rastreado else ""

    for producto in productos:
        variantes = producto.variantes
        imagenes = producto.imagenes
        total_filas = max(len(variantes), len(imagenes), 1)

        for indice in range(total_filas):
            fila = {columna: "" for columna in columnas}
            fila["Handle"] = producto.handle

            # --- campos del producto: sólo en la primera fila ---
            if indice == 0:
                fila["Title"] = producto.titulo
                fila["Body (HTML)"] = producto.body_html
                fila["Vendor"] = producto.vendor
                fila["Product Category"] = rules.PRODUCT_CATEGORY_POR_DEFECTO
                fila["Type"] = producto.tipo
                fila["Tags"] = producto.tags
                fila["Published"] = producto.publicado
                fila["Gift Card"] = defaults["Gift Card"]
                fila["Status"] = producto.estatus
                # Filtros de la tienda: sólo en la primera fila, como los
                # demás datos del producto, y sólo si la columna va en la salida.
                for columna, valor in valores_filtros(producto).items():
                    if columna in fila:
                        fila[columna] = valor
                if variantes:
                    fila["Option1 Name"] = producto.option1_name
                    if producto.option2_name:
                        fila["Option2 Name"] = producto.option2_name

            # --- campos de variante ---
            if indice < len(variantes):
                variante = variantes[indice]
                fila["Option1 Value"] = _celda_texto(variante.talla)
                if producto.option2_name:
                    fila["Option2 Value"] = _celda_texto(variante.color)
                fila["Variant SKU"] = _celda_texto(variante.sku)
                fila["Variant Barcode"] = _celda_texto(variante.barcode)
                fila["Variant Price"] = variante.precio
                fila["Variant Grams"] = defaults["Variant Grams"]
                fila["Variant Weight Unit"] = defaults["Variant Weight Unit"]
                fila["Variant Inventory Tracker"] = rastreador
                fila["Variant Inventory Qty"] = cantidad
                fila["Variant Inventory Policy"] = defaults["Variant Inventory Policy"]
                fila["Variant Fulfillment Service"] = defaults["Variant Fulfillment Service"]
                fila["Variant Requires Shipping"] = defaults["Variant Requires Shipping"]
                fila["Variant Taxable"] = defaults["Variant Taxable"]

            # --- imágenes ---
            if indice < len(imagenes):
                fila["Image Src"] = imagenes[indice]
                fila["Image Position"] = str(indice + 1)
                fila["Image Alt Text"] = rules.IMAGEN_ALT_TEXT

            filas.append(fila)

    df = pd.DataFrame(filas, columns=columnas)
    log.info("Layout Shopify: %d productos -> %d filas", len(productos), len(df))
    return df


def _dar_formato(hoja, df: pd.DataFrame, columnas_texto: list[str]) -> None:
    for celda in hoja[1]:
        celda.fill = _AZUL
        celda.font = _BLANCO_NEGRITA
        celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.freeze_panes = "A2"

    indices_texto = [i + 1 for i, c in enumerate(df.columns) if c in columnas_texto]
    for indice in indices_texto:
        letra = get_column_letter(indice)
        for celda in hoja[letra][1:]:
            celda.number_format = "@"

    anchos = {"Handle": 34, "Title": 55, "Body (HTML)": 60, "Tags": 55,
              "Image Src": 60, "Variant SKU": 16, "Variant Barcode": 18}
    for indice, columna in enumerate(df.columns, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = anchos.get(columna, 18)


def exportar_productos(df: pd.DataFrame, ruta: Path) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="products")
        _dar_formato(writer.sheets["products"], df, rules.COLUMNAS_TEXTO)
    log.info("Escrito %s (%d filas x %d columnas)", ruta.name, len(df), len(df.columns))


def exportar_csv(df: pd.DataFrame, ruta: Path) -> None:
    """CSV UTF-8, que es el formato que Shopify importa de forma nativa."""
    ruta.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(ruta, index=False, encoding="utf-8-sig")
    log.info("Escrito %s", ruta.name)


# ---------------------------------------------------------------------------
# Reporte del proceso
# ---------------------------------------------------------------------------

def construir_reporte(*, productos_totales: int, exportables: list[Producto],
                      rechazados: list[Producto], hallazgos: list[Hallazgo],
                      filas_generadas: int, variantes_leidas: int,
                      mapa_maestro: dict, mapa_variantes: dict,
                      columnas_plantilla: list[str],
                      verificacion_imagenes: dict,
                      inventario: str = "", inventario_rastreado: bool = False
                      ) -> dict[str, pd.DataFrame]:
    """Arma las hojas del reporte como DataFrames."""
    todos = exportables + rechazados
    por_codigo: dict[str, set[str]] = {}
    for hallazgo in hallazgos:
        clave = hallazgo.handle or hallazgo.articulo or "(sin producto)"
        por_codigo.setdefault(hallazgo.codigo, set()).add(clave)

    handles_con_error = {h.handle for h in hallazgos if h.nivel == "error" and h.handle}
    handles_con_aviso = {h.handle for h in hallazgos if h.nivel == "advertencia" and h.handle}
    correctos = [p for p in exportables
                 if p.handle not in handles_con_aviso and p.handle not in handles_con_error]

    def cuenta(codigo: str) -> int:
        return len(por_codigo.get(codigo, ()))

    resumen = [
        ("Fecha del proceso", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Archivo de origen", rules.ARCHIVO_ORIGEN),
        ("Estrategia de agrupación", rules.AGRUPACION),
        ("Estrategia de SKU", rules.SKU_ESTRATEGIA),
        ("", ""),
        ("Total de productos leídos", productos_totales),
        ("Productos procesados", len(todos)),
        ("Productos exportados al Excel", len(exportables)),
        ("Productos correctos (sin ninguna observación)", len(correctos)),
        ("Productos con advertencias", len(handles_con_aviso)),
        ("Productos con errores (NO exportados)", len(rechazados)),
        ("", ""),
        ("Variantes leídas del origen", variantes_leidas),
        ("Variantes exportadas", sum(len(p.variantes) for p in exportables)),
        ("Imágenes exportadas", sum(len(p.imagenes) for p in exportables)),
        ("Filas del Excel final", filas_generadas),
        ("Columnas del Excel final", len(columnas_plantilla)),
        ("", ""),
        ("SKUs duplicados", cuenta("sku_duplicado")),
        ("Productos sin imagen", cuenta("sin_imagen")),
        ("Productos sin precio", cuenta("precio_vacio") + cuenta("precio_invalido")),
        ("Productos sin categoría", cuenta("categoria_vacia")),
        ("Productos sin color", cuenta("sin_color")),
        ("Productos sin marca", cuenta("sin_marca")),
        ("Productos sin descripción", cuenta("sin_descripcion")),
        ("Productos sin variantes", cuenta("sin_variantes")),
        ("Productos no aprobados por el proveedor", cuenta("no_aprobado")),
        ("Handles con colisión resuelta", cuenta("handle_duplicado")),
        ("URLs de imagen con problema", cuenta("url_invalida")),
        ("Artículos con EAN sin ficha en el maestro", cuenta("variantes_huerfanas")),
        ("Inventario escrito en cada variante", inventario),
        ("Shopify lleva la cuenta del inventario",
         "sí, descuenta al vender" if inventario_rastreado
         else "no, la cantidad no cambia"),
        ("  (el concentrado no trae existencias reales)", ""),
    ]

    hojas = {
        "Resumen": pd.DataFrame(resumen, columns=["Concepto", "Valor"]),

        "Detalle por producto": pd.DataFrame([{
            "Handle": p.handle,
            "Nº artículo": p.articulo,
            "Estilo": p.estilo,
            "Título": p.titulo,
            "Marca": p.vendor,
            "Type": p.tipo,
            "Color": p.color,
            "Familia": p.familia,
            "Aprobación": p.aprobacion,
            "Status": p.estatus,
            "Variantes": len(p.variantes),
            "Imágenes": len(p.imagenes),
            "Precio mín.": min((v.precio for v in p.variantes if v.precio), default=""),
            "Precio máx.": max((v.precio for v in p.variantes if v.precio), default=""),
            "Tallas": ", ".join(v.talla for v in p.variantes),
            "Exportado": "Sí" if p in exportables else "NO",
        } for p in todos]),

        "Problemas": pd.DataFrame([{
            "Nivel": h.nivel.upper(),
            "Código": h.codigo,
            "Handle": h.handle,
            "Nº artículo": h.articulo,
            "Título": h.titulo,
            "Qué pasó": h.mensaje,
            "Detalle": h.detalle,
        } for h in sorted(hallazgos, key=lambda x: (x.nivel != "error", x.codigo, x.handle))]),

        "Conteo de problemas": pd.DataFrame(
            sorted(
                [{"Código": c, "Productos afectados": len(v),
                  "Nivel": rules.NIVEL_VALIDACION.get(c, "advertencia")}
                 for c, v in por_codigo.items()],
                key=lambda d: -d["Productos afectados"],
            )
        ),

        "Mapeo de columnas": pd.DataFrame(
            [{"Hoja": rules.HOJA_MAESTRO, "Campo lógico": k, "Columna encontrada": v}
             for k, v in sorted(mapa_maestro.items())]
            + [{"Hoja": " + ".join(rules.HOJAS_VARIANTES), "Campo lógico": k,
                "Columna encontrada": v} for k, v in sorted(mapa_variantes.items())]
        ),

        "Columnas sin origen": pd.DataFrame(
            [{"Columna de la plantilla": c,
              "Motivo": "El archivo del proveedor no trae este dato; se deja vacía."}
             for c in rules.COLUMNAS_SIN_ORIGEN if c in columnas_plantilla]
        ),
    }

    if verificacion_imagenes:
        hojas["Imágenes verificadas"] = pd.DataFrame([
            {"URL": r["url"], "Accesible": "Sí" if r["accesible"] else "No",
             "Código HTTP": r["codigo"], "Tipo": r["content_type"],
             "Bytes": r["bytes"], "Error": r["error"]}
            for r in verificacion_imagenes.values()
        ])

    return hojas


def exportar_reporte(hojas: dict[str, pd.DataFrame], ruta: Path) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            if df.empty:
                df = pd.DataFrame([{"": "Sin registros"}])
            df.to_excel(writer, index=False, sheet_name=nombre[:31])
            hoja = writer.sheets[nombre[:31]]
            for celda in hoja[1]:
                celda.fill = _AZUL
                celda.font = _BLANCO_NEGRITA
            hoja.freeze_panes = "A2"
            for indice, columna in enumerate(df.columns, start=1):
                largo = max([len(str(columna))] +
                            [len(str(v)) for v in df[columna].head(200)]) + 2
                hoja.column_dimensions[get_column_letter(indice)].width = min(largo, 70)
    log.info("Escrito %s (%d hojas)", ruta.name, len(hojas))
